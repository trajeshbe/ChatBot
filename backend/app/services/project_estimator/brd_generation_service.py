"""
BRD (Business Requirements Document) Generation Service

Generates professional PowerPoint BRD documents using LLM-powered content generation.
Based on analysis of real project proposals.

Author: AI Assistant
Date: 2025-11-21
"""

import logging
from typing import Dict, Any, List, Optional, Union
from pathlib import Path
import asyncio

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

logger = logging.getLogger(__name__)


class BRDGenerationService:
    """
    Generate BRD PowerPoint documents using LLM + python-pptx.

    This service generates professional consulting-style BRD proposals
    with 12 standard sections, matching the structure and tone of
    real project proposals.
    """

    def __init__(self, llm_client):
        """
        Initialize BRD generation service.

        Args:
            llm_client: LLM client (OpenAI, Anthropic, or Ollama)
        """
        self.llm_client = llm_client
        logger.info("BRDGenerationService initialized")

    async def generate_brd_content(
        self,
        project_info: Dict[str, Any],
        scope_details: Dict[str, Any],
        project_type: str,
        custom_assumptions: Optional[List[str]] = None,
        custom_benefits: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """
        Generate all BRD sections using LLM.

        Args:
            project_info: Basic project information
            scope_details: Scope, objectives, deliverables
            project_type: POC, Staff Augmentation, or Full Service
            custom_assumptions: Optional user-provided assumptions
            custom_benefits: Optional user-provided benefits

        Returns:
            Dictionary with all BRD sections as structured content
        """
        logger.info(f"Generating BRD content for {project_info.get('projectName')}")

        try:
            # Generate all sections in parallel for efficiency
            tasks = [
                self._generate_introduction(project_info, project_type),
                self._generate_objectives(project_info, scope_details),
                self._generate_scope(scope_details),
                self._generate_workflow(project_info, project_type),
                self._generate_deliverables(project_type, scope_details),
                self._generate_assumptions(project_type, custom_assumptions),
                self._generate_benefits(project_type, custom_benefits),
            ]

            results = await asyncio.gather(*tasks)

            brd_content = {
                "introduction": results[0],
                "objectives": results[1],
                "scope": results[2],
                "workflow": results[3],
                "deliverables": results[4],
                "assumptions": results[5],
                "benefits": results[6],
                "project_info": project_info
            }

            logger.info("BRD content generation complete")
            return brd_content

        except Exception as e:
            logger.error(f"Error generating BRD content: {e}")
            raise

    async def _generate_introduction(
        self,
        project_info: Dict[str, Any],
        project_type: str
    ) -> str:
        """Generate introduction section using LLM"""

        prompt = f"""
Write a professional introduction for a {project_type} project proposal.

Project Name: {project_info.get('projectName', 'Data Extraction Solution')}
Client: {project_info.get('clientName', '')}
Industry: {project_info.get('industry', '')}
Description: {project_info.get('projectDescription', '')}

Write a 2-3 paragraph introduction that:
1. Acknowledges the business challenge or opportunity
2. Introduces the proposed solution at a high level
3. Highlights the key value proposition
4. Uses professional consulting tone

Reference style:
"In today's data-driven business environment, organizations face the challenge of efficiently extracting and processing information from diverse sources. This proposal outlines an automated data extraction solution that will streamline data collection, reduce manual effort, and ensure data accuracy and reliability."

Write the introduction:
"""

        response = await self._call_llm(prompt, max_tokens=400)
        return response.strip()

    async def _generate_objectives(
        self,
        project_info: Dict[str, Any],
        scope_details: Dict[str, Any]
    ) -> List[str]:
        """Generate 3-6 project objectives using LLM"""

        objectives_from_scope = scope_details.get('objectives', [])

        if objectives_from_scope and len(objectives_from_scope) >= 3:
            # User provided objectives, just refine them
            prompt = f"""
Refine the following project objectives into professional consulting-style statements:

{chr(10).join(f'- {obj}' for obj in objectives_from_scope)}

Make them clear, specific, and action-oriented. Start with verbs like:
- "To demonstrate..."
- "To automate..."
- "To streamline..."
- "To ensure..."
- "To reduce..."
- "To enable..."

Return only the refined objectives as a numbered list.
"""
        else:
            # Generate from scratch
            prompt = f"""
Generate 3-6 project objectives for this {project_info.get('projectType', 'POC')} project.

Project: {project_info.get('projectName', '')}
Industry: {project_info.get('industry', '')}
Description: {project_info.get('projectDescription', '')}

Each objective should:
- Start with an action verb
- Be specific and measurable
- Focus on business outcomes
- Be 1-2 sentences long

Reference examples from real proposals:
1. To demonstrate capabilities to automate the end-to-end data collection and extraction process from various sources
2. Automates data extraction from structured and unstructured sources
3. Ensures accuracy through validation, deduplication, and confidence scoring
4. Reduces manual effort by 50-60% through intelligent automation
5. Provides scalable architecture supporting high-volume data processing

Generate {max(3, len(objectives_from_scope))} objectives:
"""

        response = await self._call_llm(prompt, max_tokens=600)

        # Parse into list
        objectives = []
        for line in response.split('\n'):
            line = line.strip()
            # Remove numbering if present
            if line and (line[0].isdigit() or line.startswith('-')):
                # Remove leading number/bullet
                line = line.lstrip('0123456789.-) ')
            if line and len(line) > 10:
                objectives.append(line)

        return objectives[:6]  # Limit to 6

    async def _generate_scope(
        self,
        scope_details: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        Generate scope section with in-scope and out-of-scope items.

        Returns:
            {
                "in_scope": List[str],
                "out_of_scope": List[str]
            }
        """

        in_scope = scope_details.get('inScope', [])
        out_of_scope = scope_details.get('outOfScope', [])

        if not in_scope:
            # Generate typical in-scope items
            prompt = f"""
Generate 5-8 in-scope items for a data extraction project.

Project description: {scope_details.get('projectDescription', '')}

Include typical items like:
- Data source integrations
- Extraction mechanisms
- Data processing and transformation
- Quality assurance and validation
- Documentation and handoff

Format as bullet points. Be specific.
"""
            response = await self._call_llm(prompt, max_tokens=400)
            in_scope = self._parse_bullet_list(response)

        if not out_of_scope:
            # Generate typical out-of-scope items
            prompt = """
Generate 3-5 out-of-scope items for a data extraction POC/project.

Include typical exclusions like:
- Ongoing support beyond warranty period
- Infrastructure hosting costs
- Third-party API subscription fees
- Custom feature development post-delivery
- Training and change management

Format as bullet points.
"""
            response = await self._call_llm(prompt, max_tokens=300)
            out_of_scope = self._parse_bullet_list(response)

        return {
            "in_scope": in_scope,
            "out_of_scope": out_of_scope
        }

    async def _generate_workflow(
        self,
        project_info: Dict[str, Any],
        project_type: str
    ) -> List[str]:
        """Generate workflow/process steps"""

        prompt = f"""
Generate a 5-7 step workflow for a {project_type} data extraction solution.

Project: {project_info.get('projectName', '')}

Include typical phases like:
1. Data Source Identification & Connection
2. Data Extraction & Parsing
3. Data Transformation & Normalization
4. Quality Validation & Error Handling
5. Data Storage & Delivery
6. Monitoring & Logging
7. Testing & UAT

Write each step as a clear, action-oriented statement.
Format as numbered list.
"""

        response = await self._call_llm(prompt, max_tokens=500)

        # Parse workflow steps
        steps = []
        for line in response.split('\n'):
            line = line.strip()
            if line and (line[0].isdigit() or line.startswith('-')):
                # Remove leading number/bullet
                line = line.lstrip('0123456789.-) ')
            if line and len(line) > 10:
                steps.append(line)

        return steps[:7]  # Limit to 7 steps

    async def _generate_deliverables(
        self,
        project_type: str,
        scope_details: Dict[str, Any]
    ) -> List[str]:
        """Generate project deliverables list"""

        deliverables_from_scope = scope_details.get('deliverables', [])

        if deliverables_from_scope and len(deliverables_from_scope) >= 3:
            return deliverables_from_scope

        prompt = f"""
Generate 5-8 project deliverables for a {project_type} data extraction solution.

Include typical deliverables:
- Source code and documentation
- Deployment scripts and configurations
- Data extraction pipeline (running)
- Test cases and test results
- User documentation
- Technical architecture document
- Handoff and knowledge transfer sessions

Format as bullet points. Be specific and professional.
"""

        response = await self._call_llm(prompt, max_tokens=400)
        deliverables = self._parse_bullet_list(response)

        return deliverables[:8]

    async def _generate_assumptions(
        self,
        project_type: str,
        custom_assumptions: Optional[List[str]] = None
    ) -> List[str]:
        """
        Generate project assumptions.

        Assumptions are stated as facts, not "We assume that..."
        Example: "Client will provide access to data sources within 2 business days"
        """

        if custom_assumptions and len(custom_assumptions) >= 3:
            return custom_assumptions

        prompt = f"""
Generate 5-8 project assumptions for a {project_type} data extraction project.

State each assumption as a fact, not as "We assume...".

Good examples from real proposals:
- "Client will provide access to all required data sources within 2 business days of project kickoff"
- "All third-party APIs will remain stable and accessible during the project timeline"
- "Client will provide timely feedback during UAT phase (within 48 hours)"
- "Data sources will maintain consistent structure throughout the project"
- "Client technical team will be available for integration support"

{"Focus on POC-specific assumptions: limited scope, proof-of-concept nature, no ongoing support" if project_type == "POC" else ""}
{"Focus on full service assumptions: ongoing support, maintenance, SLA commitments" if project_type == "Full Service" else ""}

Generate 5-8 assumptions:
"""

        response = await self._call_llm(prompt, max_tokens=500)
        assumptions = self._parse_bullet_list(response)

        return assumptions[:8]

    async def _generate_benefits(
        self,
        project_type: str,
        custom_benefits: Optional[List[str]] = None
    ) -> List[str]:
        """
        Generate key benefits/value adds.

        Benefits should highlight FOC items, efficiency gains, cost savings.
        """

        if custom_benefits and len(custom_benefits) >= 3:
            return custom_benefits

        foc_note = ""
        if project_type == "POC":
            foc_note = "\nFor POC, emphasize: Project Management and Documentation are Free of Charge"

        prompt = f"""
Generate 5-8 key benefits/value adds for a {project_type} data extraction solution.

Highlight:
- Efficiency improvements (50-60% reduction in manual effort)
- Cost savings
- Scalability
- Data accuracy improvements
- Time-to-market reduction
{foc_note}

Examples from real proposals:
- "Automation reduces manual data entry effort by 50-60%"
- "Project Management and Documentation Free of Charge (FOC)"
- "Scalable architecture supports 10x growth in data volume"
- "Faster time-to-market with 8-week POC timeline"
- "Improved data accuracy through automated validation"

Format as bullet points. Be specific with percentages and metrics where possible.
"""

        response = await self._call_llm(prompt, max_tokens=500)
        benefits = self._parse_bullet_list(response)

        return benefits[:8]

    def create_powerpoint(
        self,
        brd_content: Dict[str, Any],
        output_path: str
    ) -> str:
        """
        Generate PowerPoint file from BRD content.

        Creates a professional proposal with 12-14 slides following
        the structure of analyzed sample proposals.

        Args:
            brd_content: Dictionary with all BRD sections
            output_path: Path where PPTX file should be saved

        Returns:
            Path to generated PPTX file
        """
        logger.info(f"Creating PowerPoint at {output_path}")

        try:
            prs = Presentation()
            prs.slide_width = Inches(10)
            prs.slide_height = Inches(7.5)

            project_info = brd_content.get('project_info', {})

            # Slide 1: Title Slide
            self._add_title_slide(prs, project_info)

            # Slide 2: Introduction
            self._add_content_slide(
                prs,
                "Introduction",
                brd_content.get("introduction", "")
            )

            # Slide 3: Objectives
            project_type = project_info.get('projectType', 'POC')
            self._add_bullet_slide(
                prs,
                f"{project_type} Objectives",
                brd_content.get("objectives", [])
            )

            # Slide 4-5: Scope (In-Scope and Out-of-Scope)
            scope = brd_content.get("scope", {})
            self._add_bullet_slide(
                prs,
                "Scope - In Scope",
                scope.get("in_scope", [])
            )
            self._add_bullet_slide(
                prs,
                "Scope - Out of Scope",
                scope.get("out_of_scope", [])
            )

            # Slide 6-7: Workflow/Process
            workflow_steps = brd_content.get("workflow", [])
            # Split into 2 slides if more than 4 steps
            if len(workflow_steps) > 4:
                mid = len(workflow_steps) // 2
                self._add_bullet_slide(
                    prs,
                    "Proposed Workflow - Part 1",
                    workflow_steps[:mid]
                )
                self._add_bullet_slide(
                    prs,
                    "Proposed Workflow - Part 2",
                    workflow_steps[mid:]
                )
            else:
                self._add_bullet_slide(
                    prs,
                    "Proposed Workflow",
                    workflow_steps
                )

            # Slide 8: Deliverables
            self._add_bullet_slide(
                prs,
                f"{project_type} Deliverables",
                brd_content.get("deliverables", [])
            )

            # Slide 9: Assumptions
            self._add_bullet_slide(
                prs,
                "Assumptions",
                brd_content.get("assumptions", [])
            )

            # Slide 10: Benefits/Value Adds
            self._add_bullet_slide(
                prs,
                "Key Value Adds",
                brd_content.get("benefits", [])
            )

            # Slide 11: Cost Summary (placeholder)
            self._add_cost_summary_slide(prs, project_info)

            # Slide 12: Timeline (placeholder)
            self._add_timeline_slide(prs, project_info)

            # Slide 13: Thank You / Contact
            self._add_closing_slide(prs, project_info)

            # Save presentation
            Path(output_path).parent.mkdir(parents=True, exist_ok=True)
            prs.save(output_path)

            logger.info(f"PowerPoint created successfully: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error creating PowerPoint: {e}")
            raise

    def _add_title_slide(
        self,
        prs: Presentation,
        project_info: Dict[str, Any]
    ):
        """Add title slide"""
        slide_layout = prs.slide_layouts[0]  # Title slide layout
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = project_info.get('projectName', 'Project Proposal')

        if slide.placeholders[1]:
            subtitle = slide.placeholders[1]
            subtitle.text = f"""Solution Approach Note and Estimates
For {project_info.get('clientName', 'Client')}
Version {project_info.get('version', '1.0')}
{project_info.get('date', '2025')}"""

    def _add_content_slide(
        self,
        prs: Presentation,
        title: str,
        content: str
    ):
        """Add slide with title and paragraph content"""
        slide_layout = prs.slide_layouts[1]  # Title and content layout
        slide = prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = title

        # Add content as text
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        # Split content into paragraphs
        paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]

        for i, para in enumerate(paragraphs):
            if i > 0:
                text_frame.add_paragraph()
            p = text_frame.paragraphs[i]
            p.text = para
            p.font.size = Pt(14)
            p.space_after = Pt(12)

    def _add_bullet_slide(
        self,
        prs: Presentation,
        title: str,
        bullets: List[str]
    ):
        """Add slide with title and bullet points"""
        slide_layout = prs.slide_layouts[1]  # Title and content layout
        slide = prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = title

        # Add bullets
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        for i, bullet_text in enumerate(bullets):
            if i > 0:
                text_frame.add_paragraph()
            p = text_frame.paragraphs[i]
            p.text = bullet_text
            p.level = 0
            p.font.size = Pt(14)

    def _add_cost_summary_slide(
        self,
        prs: Presentation,
        project_info: Dict[str, Any]
    ):
        """Add cost summary placeholder slide"""
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = "Cost Summary"

        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        p = text_frame.paragraphs[0]
        p.text = "Detailed cost breakdown provided in accompanying Excel file"
        p.font.size = Pt(16)
        p.alignment = PP_ALIGN.CENTER

    def _add_timeline_slide(
        self,
        prs: Presentation,
        project_info: Dict[str, Any]
    ):
        """Add timeline placeholder slide"""
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = "Project Timeline"

        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()

        timeline_weeks = project_info.get('timelineDuration', 8)
        project_type = project_info.get('projectType', 'POC')

        milestones = [
            f"Total Duration: {timeline_weeks} weeks",
            f"Project Type: {project_type}",
            "Week 1-2: Planning and Setup",
            f"Week 3-{timeline_weeks-2}: Development and Integration",
            f"Week {timeline_weeks-1}: Testing and UAT",
            f"Week {timeline_weeks}: Deployment and Handoff"
        ]

        for i, milestone in enumerate(milestones):
            if i > 0:
                text_frame.add_paragraph()
            p = text_frame.paragraphs[i]
            p.text = milestone
            p.font.size = Pt(14)
            if i == 0:
                p.font.bold = True

    def _add_closing_slide(
        self,
        prs: Presentation,
        project_info: Dict[str, Any]
    ):
        """Add closing/thank you slide"""
        slide_layout = prs.slide_layouts[0]  # Title layout
        slide = prs.slides.add_slide(slide_layout)

        title = slide.shapes.title
        title.text = "Thank You"

        if slide.placeholders[1]:
            subtitle = slide.placeholders[1]
            subtitle.text = f"""Questions?

Project: {project_info.get('projectName', '')}
Client: {project_info.get('clientName', '')}"""

    async def _call_llm(
        self,
        prompt: str,
        max_tokens: int = 500,
        temperature: float = 0.7
    ) -> str:
        """
        Call LLM with prompt and return response.

        Abstracts LLM provider (OpenAI, Anthropic, Ollama).
        """
        try:
            # The actual LLM client call will depend on the interface
            # This is a placeholder that should match your LLM service
            response = await self.llm_client.generate(
                prompt=prompt,
                max_tokens=max_tokens,
                temperature=temperature
            )

            return response

        except Exception as e:
            logger.error(f"LLM call failed: {e}")
            raise

    def _parse_bullet_list(self, text: str) -> List[str]:
        """Parse LLM response into clean bullet list"""
        bullets = []

        for line in text.split('\n'):
            line = line.strip()

            # Remove leading bullets/numbers
            if line and (line[0].isdigit() or line.startswith(('-', '•', '*', '○'))):
                line = line.lstrip('0123456789.-•*○) ')

            # Only keep substantial lines
            if line and len(line) > 10 and not line.startswith('#'):
                bullets.append(line)

        return bullets
