#!/usr/bin/env python3
"""
Screener.in Data Extraction Diagnostics and Debugging Script

This script extracts financial data from screener.in company pages and exports to Excel.
Includes comprehensive diagnostics and error handling.

Usage:
    python diagnose-screener-extraction.py [url] [output_file]

Examples:
    python diagnose-screener-extraction.py https://www.screener.in/company/BHARTIARTL/consolidated/
    python diagnose-screener-extraction.py https://www.screener.in/company/BHARTIARTL/consolidated/ bharti_data.xlsx

Features:
    - Handles JavaScript-rendered content with Playwright
    - Extracts multiple tables (financials, ratios, quarterly results, etc.)
    - Exports to multi-sheet Excel workbook
    - Comprehensive diagnostics and logging
    - Error recovery and retry logic
    - Visual HTML debugging
"""

import asyncio
import sys
import os
import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Any
from urllib.parse import urlparse

# Add backend to path
sys.path.insert(0, str(Path(__file__).parent.parent.parent / 'backend'))

try:
    from playwright.async_api import async_playwright, Page, Browser
    import pandas as pd
    from bs4 import BeautifulSoup
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError as e:
    print(f"❌ Missing required package: {e}")
    print("\nInstall dependencies:")
    print("  pip install playwright pandas beautifulsoup4 lxml openpyxl")
    print("  playwright install chromium")
    sys.exit(1)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler('screener_extraction.log')
    ]
)
logger = logging.getLogger(__name__)


class ScreenerExtractor:
    """Extract financial data from screener.in"""

    def __init__(self, url: str, output_file: str = None):
        self.url = url
        self.output_file = output_file or self._generate_output_filename()
        self.browser: Optional[Browser] = None
        self.page: Optional[Page] = None
        self.html_content: str = ""
        self.extracted_data: Dict[str, Any] = {}

    def _generate_output_filename(self) -> str:
        """Generate output filename from URL"""
        parsed = urlparse(self.url)
        company_path = parsed.path.strip('/').split('/')
        company_code = company_path[1] if len(company_path) > 1 else 'screener'
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"screener_{company_code}_{timestamp}.xlsx"

    async def initialize_browser(self) -> None:
        """Initialize Playwright browser with realistic settings"""
        logger.info("🚀 Initializing Playwright browser...")

        playwright = await async_playwright().start()

        # Launch browser with realistic settings to avoid detection
        self.browser = await playwright.chromium.launch(
            headless=True,  # Set to False to see browser for debugging
            args=[
                '--disable-blink-features=AutomationControlled',
                '--disable-dev-shm-usage',
                '--no-sandbox',
                '--disable-setuid-sandbox',
                '--disable-web-security'
            ]
        )

        # Create context with realistic user agent
        context = await self.browser.new_context(
            user_agent='Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
            viewport={'width': 1920, 'height': 1080},
            java_script_enabled=True,
            accept_downloads=True
        )

        self.page = await context.new_page()
        logger.info("✅ Browser initialized successfully")

    async def fetch_page(self, max_retries: int = 3) -> bool:
        """Fetch page with retry logic"""
        for attempt in range(max_retries):
            try:
                logger.info(f"📡 Fetching URL (attempt {attempt + 1}/{max_retries}): {self.url}")

                # Navigate to page
                response = await self.page.goto(
                    self.url,
                    wait_until='networkidle',
                    timeout=30000
                )

                if response.status != 200:
                    logger.warning(f"⚠️  HTTP {response.status} received")
                    if attempt < max_retries - 1:
                        await asyncio.sleep(2 ** attempt)
                        continue
                    return False

                # Wait for content to load
                logger.info("⏳ Waiting for content to load...")

                # Wait for key elements (adjust selectors as needed)
                try:
                    await self.page.wait_for_selector('table', timeout=10000)
                    logger.info("✅ Tables found on page")
                except Exception as e:
                    logger.warning(f"⚠️  Timeout waiting for tables: {e}")

                # Additional wait for JavaScript to render
                await asyncio.sleep(2)

                # Get page content
                self.html_content = await self.page.content()
                logger.info(f"✅ Fetched {len(self.html_content):,} bytes of HTML")

                # Save HTML for debugging
                debug_file = self.output_file.replace('.xlsx', '_debug.html')
                with open(debug_file, 'w', encoding='utf-8') as f:
                    f.write(self.html_content)
                logger.info(f"💾 Saved HTML debug file: {debug_file}")

                return True

            except Exception as e:
                logger.error(f"❌ Error fetching page (attempt {attempt + 1}): {e}")
                if attempt < max_retries - 1:
                    await asyncio.sleep(2 ** attempt)
                else:
                    return False

        return False

    def parse_html(self) -> None:
        """Parse HTML and extract data"""
        logger.info("🔍 Parsing HTML content...")

        soup = BeautifulSoup(self.html_content, 'lxml')

        # Extract page metadata
        self.extracted_data['metadata'] = self._extract_metadata(soup)

        # Extract all tables
        self.extracted_data['tables'] = self._extract_tables(soup)

        # Extract company info
        self.extracted_data['company_info'] = self._extract_company_info(soup)

        # Extract financial highlights
        self.extracted_data['highlights'] = self._extract_highlights(soup)

        logger.info(f"✅ Extracted {len(self.extracted_data['tables'])} tables")
        logger.info(f"✅ Extracted company info: {self.extracted_data['company_info'].get('name', 'N/A')}")

    def _extract_metadata(self, soup: BeautifulSoup) -> Dict[str, str]:
        """Extract page metadata"""
        title = soup.find('title')
        title_text = title.string if title else "Unknown"

        meta_desc = soup.find('meta', attrs={'name': 'description'})
        description = meta_desc['content'] if meta_desc else ""

        return {
            'title': title_text,
            'description': description,
            'url': self.url,
            'extracted_at': datetime.now().isoformat()
        }

    def _extract_company_info(self, soup: BeautifulSoup) -> Dict[str, str]:
        """Extract company information"""
        info = {}

        # Company name
        name_elem = soup.find('h1', class_='h2')
        if name_elem:
            info['name'] = name_elem.get_text(strip=True)

        # Company details from various sections
        # Adjust selectors based on actual page structure
        details = soup.find_all('li', class_='flex')
        for detail in details:
            spans = detail.find_all('span')
            if len(spans) >= 2:
                key = spans[0].get_text(strip=True).rstrip(':')
                value = spans[1].get_text(strip=True)
                info[key] = value

        return info

    def _extract_highlights(self, soup: BeautifulSoup) -> Dict[str, str]:
        """Extract financial highlights/key metrics"""
        highlights = {}

        # Look for highlight sections (adjust selectors as needed)
        highlight_section = soup.find('div', id='top-ratios') or soup.find('section', id='ratios')

        if highlight_section:
            items = highlight_section.find_all('li')
            for item in items:
                # Extract name and value
                name_elem = item.find('span', class_='name')
                value_elem = item.find('span', class_='number')

                if name_elem and value_elem:
                    highlights[name_elem.get_text(strip=True)] = value_elem.get_text(strip=True)

        return highlights

    def _extract_tables(self, soup: BeautifulSoup) -> Dict[str, pd.DataFrame]:
        """Extract all tables from the page"""
        tables = {}

        # Find all tables
        html_tables = soup.find_all('table')
        logger.info(f"📊 Found {len(html_tables)} tables")

        for idx, table in enumerate(html_tables):
            try:
                # Try to identify table by header or surrounding context
                table_name = self._identify_table(table, idx)

                # Parse table to DataFrame
                df = self._parse_table_to_dataframe(table)

                if df is not None and not df.empty:
                    tables[table_name] = df
                    logger.info(f"  ✅ Table '{table_name}': {df.shape[0]} rows × {df.shape[1]} columns")
                else:
                    logger.warning(f"  ⚠️  Table {idx} is empty or couldn't be parsed")

            except Exception as e:
                logger.error(f"  ❌ Error parsing table {idx}: {e}")

        return tables

    def _identify_table(self, table: BeautifulSoup, default_idx: int) -> str:
        """Identify table name from context"""
        # Try to find heading before table
        prev = table.find_previous(['h2', 'h3', 'h4'])
        if prev:
            name = prev.get_text(strip=True)
            if name and len(name) < 50:
                return self._sanitize_sheet_name(name)

        # Try to find section id
        section = table.find_parent(['section', 'div'], id=True)
        if section and section.get('id'):
            return self._sanitize_sheet_name(section['id'])

        # Check if table has a caption
        caption = table.find('caption')
        if caption:
            return self._sanitize_sheet_name(caption.get_text(strip=True))

        # Default name
        return f"Table_{default_idx + 1}"

    def _sanitize_sheet_name(self, name: str) -> str:
        """Sanitize name for Excel sheet (max 31 chars, no special chars)"""
        # Remove invalid characters
        invalid_chars = [':', '\\', '/', '?', '*', '[', ']']
        for char in invalid_chars:
            name = name.replace(char, '_')

        # Truncate to 31 characters
        return name[:31]

    def _parse_table_to_dataframe(self, table: BeautifulSoup) -> Optional[pd.DataFrame]:
        """Parse HTML table to pandas DataFrame"""
        try:
            # Extract headers
            headers = []
            header_row = table.find('thead')
            if header_row:
                header_cells = header_row.find_all(['th', 'td'])
                headers = [cell.get_text(strip=True) for cell in header_cells]
            else:
                # Try first row as header
                first_row = table.find('tr')
                if first_row:
                    header_cells = first_row.find_all(['th', 'td'])
                    headers = [cell.get_text(strip=True) for cell in header_cells]

            # Extract data rows
            rows = []
            tbody = table.find('tbody') or table
            data_rows = tbody.find_all('tr')

            # Skip header row if we already processed it
            if not table.find('thead') and headers:
                data_rows = data_rows[1:]

            for row in data_rows:
                cells = row.find_all(['td', 'th'])
                row_data = [cell.get_text(strip=True) for cell in cells]
                if row_data:  # Skip empty rows
                    rows.append(row_data)

            if not rows:
                return None

            # Create DataFrame
            if headers and len(headers) > 0:
                # Ensure all rows have same number of columns
                max_cols = max(len(headers), max(len(row) for row in rows))
                headers = headers + [''] * (max_cols - len(headers))
                rows = [row + [''] * (max_cols - len(row)) for row in rows]

                df = pd.DataFrame(rows, columns=headers[:len(rows[0])])
            else:
                df = pd.DataFrame(rows)

            return df

        except Exception as e:
            logger.error(f"Error parsing table: {e}")
            return None

    def export_to_excel(self) -> bool:
        """Export extracted data to Excel"""
        logger.info(f"📝 Exporting data to Excel: {self.output_file}")

        try:
            with pd.ExcelWriter(self.output_file, engine='openpyxl') as writer:
                # Write metadata sheet
                if self.extracted_data.get('metadata'):
                    metadata_df = pd.DataFrame([self.extracted_data['metadata']])
                    metadata_df.to_excel(writer, sheet_name='Metadata', index=False)
                    logger.info("  ✅ Wrote Metadata sheet")

                # Write company info sheet
                if self.extracted_data.get('company_info'):
                    info_df = pd.DataFrame([
                        {'Field': k, 'Value': v}
                        for k, v in self.extracted_data['company_info'].items()
                    ])
                    info_df.to_excel(writer, sheet_name='Company Info', index=False)
                    logger.info("  ✅ Wrote Company Info sheet")

                # Write highlights sheet
                if self.extracted_data.get('highlights'):
                    highlights_df = pd.DataFrame([
                        {'Metric': k, 'Value': v}
                        for k, v in self.extracted_data['highlights'].items()
                    ])
                    highlights_df.to_excel(writer, sheet_name='Highlights', index=False)
                    logger.info("  ✅ Wrote Highlights sheet")

                # Write table sheets
                for table_name, df in self.extracted_data.get('tables', {}).items():
                    df.to_excel(writer, sheet_name=table_name, index=False)
                    logger.info(f"  ✅ Wrote sheet: {table_name}")

            # Apply formatting
            self._apply_excel_formatting()

            logger.info(f"✅ Successfully exported to {self.output_file}")
            return True

        except Exception as e:
            logger.error(f"❌ Error exporting to Excel: {e}")
            return False

    def _apply_excel_formatting(self) -> None:
        """Apply formatting to Excel workbook"""
        try:
            wb = openpyxl.load_workbook(self.output_file)

            # Format each sheet
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]

                # Format header row
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Auto-adjust column widths
                for column in ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter

                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass

                    adjusted_width = min(max_length + 2, 50)
                    ws.column_dimensions[column_letter].width = adjusted_width

            wb.save(self.output_file)
            logger.info("✅ Applied Excel formatting")

        except Exception as e:
            logger.error(f"⚠️  Error applying formatting: {e}")

    async def cleanup(self) -> None:
        """Cleanup resources"""
        if self.browser:
            await self.browser.close()
            logger.info("🧹 Browser closed")

    async def extract(self) -> bool:
        """Main extraction workflow"""
        try:
            # Initialize browser
            await self.initialize_browser()

            # Fetch page
            if not await self.fetch_page():
                logger.error("❌ Failed to fetch page")
                return False

            # Parse HTML
            self.parse_html()

            # Export to Excel
            if not self.export_to_excel():
                logger.error("❌ Failed to export to Excel")
                return False

            # Print summary
            self._print_summary()

            return True

        except Exception as e:
            logger.error(f"❌ Extraction failed: {e}", exc_info=True)
            return False

        finally:
            await self.cleanup()

    def _print_summary(self) -> None:
        """Print extraction summary"""
        print("\n" + "="*80)
        print("📊 EXTRACTION SUMMARY")
        print("="*80)
        print(f"URL: {self.url}")
        print(f"Output File: {self.output_file}")
        print(f"Company: {self.extracted_data.get('company_info', {}).get('name', 'N/A')}")
        print(f"\nExtracted Data:")
        print(f"  • Tables: {len(self.extracted_data.get('tables', {}))}")
        print(f"  • Company Info Fields: {len(self.extracted_data.get('company_info', {}))}")
        print(f"  • Highlights: {len(self.extracted_data.get('highlights', {}))}")

        if self.extracted_data.get('tables'):
            print(f"\nTable Details:")
            for name, df in self.extracted_data['tables'].items():
                print(f"  • {name}: {df.shape[0]} rows × {df.shape[1]} columns")

        print("\n" + "="*80)
        print(f"✅ Data exported successfully to: {self.output_file}")
        print("="*80 + "\n")


async def main():
    """Main function"""
    # Parse command line arguments
    url = sys.argv[1] if len(sys.argv) > 1 else "https://www.screener.in/company/BHARTIARTL/consolidated/"
    output_file = sys.argv[2] if len(sys.argv) > 2 else None

    print("\n" + "="*80)
    print("🔍 SCREENER.IN DATA EXTRACTION DIAGNOSTICS")
    print("="*80)
    print(f"URL: {url}")
    print(f"Output: {output_file or 'Auto-generated'}")
    print("="*80 + "\n")

    # Create extractor
    extractor = ScreenerExtractor(url, output_file)

    # Run extraction
    success = await extractor.extract()

    if success:
        print("\n✅ Extraction completed successfully!")
        sys.exit(0)
    else:
        print("\n❌ Extraction failed. Check logs for details.")
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
